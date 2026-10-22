"""Optional local DOCX -> LibreOffice PDF -> page images. No network access."""
from importlib.metadata import version
from pathlib import Path
import shutil
import subprocess
import threading
import unicodedata

from .workflow import WorkflowError
from .metrics import Trace

PDFIUM_LOCK=threading.Lock()
FONT='DejaVu Sans'
MAX_PAGES=10


def content_lines(content, markdown=False):
    lines=[(content['name'],'title'),(content['headline'],'headline')]
    lines.extend((value,'contact') for value in content['contacts'])
    if content.get('schema_version')=='2.0':
        for section in content['sections']:
            lines.append((section['heading'],'heading'))
            for group in section['groups']:
                if group['context']:lines.append((group['context'],'context'))
                lines.extend((s['text'],'bullet') for s in group['statements'])
    else:
        lines.append((content['section'],'heading'))
        lines.extend((claim,'bullet') for claim in content['claims'])
    if markdown:
        prefixes={'title':'# ','heading':'## ','context':'### ','bullet':'- '}
        return [(value,prefixes.get(kind,'')) for value,kind in lines]
    return lines


def paragraphs(content):
    return [('- ' if kind=='bullet' else '')+value for value,kind in content_lines(content)]


def build_docx(content, path):
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor
    document=Document()
    section=document.sections[0]
    section.page_width=Mm(210);section.page_height=Mm(297)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Mm(20)
    for style_name,size in [('Normal',11),('Title',21),('Heading 1',13),('Heading 2',11)]:
        style=document.styles[style_name]
        style.font.name=FONT;style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_after=Pt(7)
        style.paragraph_format.line_spacing=1.15
    document.core_properties.author=''
    document.core_properties.last_modified_by=''
    document.core_properties.title=content['name']
    document.core_properties.subject=''
    document.core_properties.comments=''
    for value,kind in content_lines(content):
        style={'title':'Title','heading':'Heading 1','context':'Heading 2'}.get(kind,'Normal')
        document.add_paragraph(('- ' if kind=='bullet' else '')+value,style)
    # The library template may carry inherited blue Title rules. Whitespace,
    # rather than theme-dependent borders, separates this simple CV layout.
    for element in [document.styles.element,document.element]:
        for border in element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    document.save(path)
    reopened=Document(path)
    if [p.text for p in reopened.paragraphs]!=paragraphs(content):
        raise WorkflowError('failed','DOCX text differs from the reviewed content model')


def normalized(value):
    return ''.join(unicodedata.normalize('NFC',value).split())


def verify_pdf(content, path):
    from pypdf import PdfReader
    reader=PdfReader(path)
    if reader.is_encrypted or not 1<=len(reader.pages)<=MAX_PAGES:
        raise WorkflowError('blocked','PDF is encrypted, empty or exceeds the 10-page CV limit')
    actual=''.join(page.extract_text() or '' for page in reader.pages)
    if normalized(actual)!=normalized(''.join(paragraphs(content))):
        raise WorkflowError('blocked','PDF text differs from content; inspect missing glyphs, clipping or conversion changes')
    for page in reader.pages:
        resources=page['/Resources'] if '/Resources' in page else {}
        fonts=resources['/Font'] if '/Font' in resources else {}
        for reference in fonts.values():
            font=reference.get_object()
            if 'DejaVuSans' not in str(font.get('/BaseFont','')).replace('-',''):
                raise WorkflowError('blocked','PDF font substitution detected; install the documented DejaVu Sans family')
            descendant=font.get('/DescendantFonts')
            descriptor=(descendant[0].get_object() if descendant else font).get('/FontDescriptor')
            if not descriptor or not any(key in descriptor.get_object() for key in ('/FontFile','/FontFile2','/FontFile3')):
                raise WorkflowError('blocked','PDF font is not embedded; do not approve this conversion')
    return len(reader.pages)


def preview_pdf(path,directory,count):
    import pypdfium2 as pdfium
    # PDFium is not thread safe, including across different document objects.
    with PDFIUM_LOCK:
        pdf=pdfium.PdfDocument(path)
        try:
            if len(pdf)!=count:raise WorkflowError('failed','PDF parsers disagree about the page count')
            for index in range(count):
                page=pdf[index]
                try:
                    bitmap=page.render(scale=2)
                    try:
                        picture=bitmap.to_pil()
                        try:picture.save(directory/f'page-{index+1:03}.png')
                        finally:picture.close()
                    finally:bitmap.close()
                finally:page.close()
        finally:pdf.close()


def render_files(content, directory: Path, soffice=None):
    trace=Trace()
    try:
        import docx
        import pypdf
        import pypdfium2 as pdfium
        import PIL
    except ImportError as error:
        raise WorkflowError('manual_required','Optional document dependencies missing; install the reviewed requirements-documents.lock in a virtual environment') from error
    executable=Path(soffice).expanduser() if soffice else Path(shutil.which('soffice') or '__missing_soffice__')
    if soffice and not executable.is_absolute():
        raise WorkflowError('needs_input','--soffice must be an explicitly trusted absolute executable path')
    if not executable.is_file():
        raise WorkflowError('manual_required','LibreOffice is unavailable; supply a reviewed local --soffice path. No render or approval was recorded')
    docx_path=directory/'cv.docx'
    with trace.stage('docx_generation_and_reopen'):build_docx(content,docx_path)
    profile=directory/'lo-profile'
    try:
        with trace.stage('converter_inventory_startup'):
            inventory=subprocess.run([str(executable),'--version'],capture_output=True,text=True,timeout=10,check=True)
        with trace.stage('converter_startup_and_conversion'):
            result=subprocess.run([str(executable),'-env:UserInstallation='+profile.resolve().as_uri(),
                '--headless','--norestore','--nodefault','--nofirststartwizard',
                '--convert-to','pdf:writer_pdf_Export','--outdir',str(directory),str(docx_path)],
                cwd=directory,capture_output=True,text=True,timeout=90,check=False)
    except (OSError,subprocess.SubprocessError) as error:
        raise WorkflowError('failed','Local document conversion failed or timed out; no automatic retry or approval') from error
    if result.returncode!=0 or not (directory/'cv.pdf').is_file():
        raise WorkflowError('failed','LibreOffice did not produce a PDF; inspect local converter installation and font availability')
    with trace.stage('pdf_text_and_embedded_font_checks'):count=verify_pdf(content,directory/'cv.pdf')
    with trace.stage('page_images'):preview_pdf(directory/'cv.pdf',directory,count)
    # Only this freshly created, private converter profile is removed.
    if profile.exists():shutil.rmtree(profile)
    return {'page_count':count,'metrics':trace.report(render_count=1,converter_processes=2,page_images=count),
            'renderer':{'libreoffice':inventory.stdout.strip(),
            'python_docx':version('python-docx'),'pypdf':version('pypdf'),
            'pypdfium2':version('pypdfium2'),'pillow':version('Pillow'),'font':FONT,'preview_dpi':144}}
